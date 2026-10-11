"""Parsing tests: all formats, invalid and empty documents."""
from __future__ import annotations

from pathlib import Path

import pytest

from app.parsing import ParsedDocument, ParseError, parse_document


def _write(tmp_path: Path, name: str, data: bytes | str) -> Path:
    p = tmp_path / name
    p.write_bytes(data if isinstance(data, bytes) else data.encode("utf-8"))
    return p


def test_parse_txt(tmp_path):
    p = _write(tmp_path, "a.txt", b"Hello\n\nWorld")
    doc = parse_document(p)
    assert "Hello" in doc.text and doc.page_count == 0


def test_parse_empty_txt(tmp_path):
    p = _write(tmp_path, "e.txt", b"")
    doc = parse_document(p)
    assert doc.text == ""


def test_parse_markdown(tmp_path):
    p = _write(tmp_path, "a.md", "# Title\n\nSome body text here.")
    doc = parse_document(p)
    assert "Title" in doc.text


def test_parse_html_strips_scripts(tmp_path):
    p = _write(
        tmp_path,
        "a.html",
        b"<html><head><title>T</title><script>alert(1)</script></head>"
        b"<body><h1>Head</h1><p>Para one.</p></body></html>",
    )
    doc = parse_document(p)
    assert "Head" in doc.text and "alert" not in doc.text
    assert doc.metadata.get("title") == "T"


def test_parse_csv(tmp_path):
    p = _write(tmp_path, "a.csv", "uid,title\nCTL-1,Do the thing\n\n")
    doc = parse_document(p)
    assert doc.metadata["row_count"] == 1
    assert "Do the thing" in doc.text


def test_parse_json(tmp_path):
    p = _write(tmp_path, "a.json", b'{"a": 1, "b": {"c": "x"}}')
    doc = parse_document(p)
    assert doc.metadata["key_count"] == 2
    assert '"x"' in doc.text or "x" in doc.text


def test_parse_invalid_json(tmp_path):
    p = _write(tmp_path, "a.json", b"{not json")
    with pytest.raises(ParseError):
        parse_document(p)


def test_parse_unsupported_type(tmp_path):
    p = _write(tmp_path, "a.exe", b"MZ")
    with pytest.raises(ParseError):
        parse_document(p)


def test_parse_docx(tmp_path):
    from docx import Document as Docx

    p = tmp_path / "a.docx"
    d = Docx()
    d.add_heading("Policy", level=1)
    d.add_paragraph("The bank shall retain records.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "k"
    table.rows[0].cells[1].text = "v"
    d.save(str(p))
    doc = parse_document(p)
    assert "retain records" in doc.text
    assert any(b.kind == "heading" for b in doc.blocks)
    assert any(b.kind == "table" for b in doc.blocks)


def test_parse_real_rbi_pdf():
    from app.parsing import parse_document as _parse

    repo = Path(__file__).resolve().parents[2]
    pdf = repo / "data" / "samples" / "official_rbi" / "RBI_Internal_Ombudsman_Directions_2023.pdf"
    if not pdf.exists():
        pytest.skip("RBI sample PDF not present")
    doc = _parse(pdf)
    assert doc.page_count >= 10
    assert "Internal Ombudsman" in doc.text
    assert len(doc.text) > 5000


def test_parse_corrupt_pdf(tmp_path):
    p = _write(tmp_path, "a.pdf", b"%PDF-1.4 garbage \x00\x01")
    with pytest.raises(ParseError):
        parse_document(p)


def test_binary_rejected_as_text():
    from app.security import SecurityError, validate_upload

    with pytest.raises(SecurityError):
        validate_upload("x.txt", "text/plain", 10, b"\x00\x01binary")
